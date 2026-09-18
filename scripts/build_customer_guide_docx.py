#!/usr/bin/env python3
"""Build the customer-facing Word guide from its maintained Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "递归智能报告系统_客户使用指南.md"
TARGET = ROOT / "docs" / "递归智能报告系统_客户使用指南.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)


def add_inline(paragraph, text: str) -> None:
    """Render the small Markdown subset used by the guide."""
    token = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\[[^]]+\]\([^)]+\))")
    cursor = 0
    for match in token.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        value = match.group(0)
        if value.startswith("**"):
            run = paragraph.add_run(value[2:-2])
            run.bold = True
        elif value.startswith("`"):
            run = paragraph.add_run(value[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(31, 78, 121)
        else:
            paragraph.add_run(value[1 : value.index("]")])
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def parse_table(lines: list[str]) -> list[list[str]]:
    return [[cell.strip() for cell in line.strip().strip("|").split("|")] for line in lines]


def is_separator_row(row: list[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)


def build() -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()

    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.3)
    section.right_margin = Cm(2.3)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    for name, size, color in [
        ("Title", 24, "17365D"),
        ("Heading 1", 17, "17365D"),
        ("Heading 2", 14, "1F4E79"),
        ("Heading 3", 12, "2F5597"),
    ]:
        style = styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True

    i = 0
    in_code = False
    code_lines: list[str] = []
    first_heading = True
    while i < len(lines):
        raw = lines[i]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                p = document.add_paragraph()
                p.style = styles["Normal"]
                p.paragraph_format.left_indent = Cm(0.8)
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after = Pt(8)
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(31, 78, 121)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(raw)
            i += 1
            continue
        if not stripped or stripped == "---":
            i += 1
            continue
        if stripped.startswith("|") and i + 1 < len(lines) and lines[i + 1].strip().startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i])
                i += 1
            rows = parse_table(table_lines)
            if len(rows) > 1 and is_separator_row(rows[1]):
                rows.pop(1)
            table = document.add_table(rows=len(rows), cols=max(map(len, rows)))
            table.style = "Table Grid"
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.autofit = True
            for r_index, row in enumerate(rows):
                for c_index, value in enumerate(row):
                    cell = table.cell(r_index, c_index)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    p = cell.paragraphs[0]
                    add_inline(p, value)
                    if r_index == 0:
                        set_cell_shading(cell, "D9EAF7")
                        for run in p.runs:
                            run.bold = True
                if r_index == 0:
                    set_repeat_table_header(table.rows[0])
            document.add_paragraph().paragraph_format.space_after = Pt(1)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1))
            title = heading.group(2)
            if level == 1 and first_heading:
                p = document.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                add_inline(p, title)
                first_heading = False
            else:
                document.add_heading(title, level=level)
            i += 1
            continue
        if stripped.startswith(">"):
            p = document.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.right_indent = Cm(0.5)
            add_inline(p, stripped.lstrip("> "))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(89, 89, 89)
            i += 1
            continue
        bullet = re.match(r"^-\s+(.+)$", stripped)
        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        checkbox = re.match(r"^-\s+\[([ xX])\]\s+(.+)$", stripped)
        if checkbox:
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, ("☑ " if checkbox.group(1).lower() == "x" else "☐ ") + checkbox.group(2))
        elif bullet:
            p = document.add_paragraph(style="List Bullet")
            add_inline(p, bullet.group(1))
        elif numbered:
            p = document.add_paragraph(style="List Number")
            add_inline(p, numbered.group(1))
        else:
            p = document.add_paragraph()
            add_inline(p, stripped.rstrip("  "))
        i += 1

    core = document.core_properties
    core.title = "递归智能报告系统客户使用指南"
    core.subject = "系统完整功能与操作说明"
    core.author = "递归智能报告系统"
    document.save(TARGET)


if __name__ == "__main__":
    build()
